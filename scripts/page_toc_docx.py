# 给校长的页面目录：python scripts/page_toc_scan.py site.json && python scripts/page_toc_docx.py site.json 出档.docx
import json, sys
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

d = json.load(open(sys.argv[1], encoding='utf-8'))
out = sys.argv[2]

NAME = {'index.html': '首页', 'about/index.html': '认识南华', 'history/index.html': '南华春秋（校史）',
        'organization/index.html': '学校组织', 'board/index.html': '董事会', 'academic/index.html': '学术课程',
        'cocurricular/index.html': '联课活动', 'facilities/index.html': '硬体设施', 'faculty/index.html': '师资团队',
        'museum/index.html': '南华数码馆', 'news/index.html': '每日南华（新闻馆）', '90/index.html': '九十周年校庆',
        'admission/index.html': '招生专区'}
ORDER = list(NAME)
PATH = lambda p: '/' if p == 'index.html' else '/' + p.replace('index.html', '')

# 首页上每个子页的入口位置（依 index.html 实际版面）
ENTRY = {'about/index.html': '顶栏、目录', 'academic/index.html': '顶栏、目录', 'cocurricular/index.html': '顶栏、目录',
         'facilities/index.html': '顶栏、目录', 'museum/index.html': '顶栏、目录、第柒节', '90/index.html': '顶栏（90 周年徽标）、目录、第柒节',
         'admission/index.html': '顶栏「招生报名」、目录、第捌节、底部快捷钮', 'history/index.html': '目录、第柒节',
         'news/index.html': '目录、第柒节、校园动态', 'organization/index.html': '目录', 'board/index.html': '目录',
         'faculty/index.html': '目录'}

NEWS = [('2018', 137), ('2019', 209), ('2020', 134), ('2021', 101), ('2022', 135), ('2023', 171), ('2024', 112), ('2025', 130), ('2026', 122)]

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Microsoft YaHei'; st.font.size = Pt(10.5)
st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
for s in ('Heading 1', 'Heading 2'):
    f = doc.styles[s].font; f.name = 'Microsoft YaHei'; f.color.rgb = None
    doc.styles[s].element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')

def table(head, rows):
    t = doc.add_table(rows=1, cols=len(head)); t.style = 'Table Grid'
    for c, h in zip(t.rows[0].cells, head): c.text = h
    for r in rows:
        for c, v in zip(t.add_row().cells, r): c.text = v

doc.add_heading('南华独中新官网　页面目录', level=1)
doc.add_paragraph('版本：改版预览版，整理日期 2026-10-23。')
doc.add_paragraph('「最后更新」指该页面最近一次修改的日期。「网址」是接在网站地址后面的路径，例如网站地址加上 /about/ 即「认识南华」。')

if len(sys.argv) > 3:  # 第三个参数：架构图 PNG（由 Mermaid 渲染）
    from docx.shared import Cm
    doc.add_heading('网站架构图', level=2)
    doc.add_paragraph('首页通往四组页面；同组内的页面并列，彼此之间另有互链（见第二节）。')
    doc.add_picture(sys.argv[3], width=Cm(16))
doc.add_heading('一、首页可以前往的页面', level=2)
doc.add_paragraph('首页本身包含：首屏影片、动画《迎领之路》、壹至捌各节、家长布告栏（最新动态、本月行事历）、校园动态、董事长与校长寄语、联系方式。')
table(['子页', '网址', '在首页的入口位置'], [(NAME[p], PATH(p), ENTRY[p]) for p in ORDER[1:]])

doc.add_heading('二、各页一览', level=2)
rows = []
for i, p in enumerate(ORDER, 1):
    v = d[p]
    goes = '、'.join(NAME[t] for t in ORDER if t in v['internal']) or '—'
    if v['news']: goes += '；另链接 %d 篇新闻' % v['news']
    if p == 'news/index.html': goes = '首页、2018–2026 各年份目录（见第三节）'
    rows.append((str(i), NAME[p], PATH(p), v['date'], goes))
table(['序', '页面', '网址', '最后更新', '从这一页可以前往'], rows)

doc.add_heading('三、每日南华（新闻馆）按年份', level=2)
doc.add_paragraph('新闻馆首页 /news/ 按年份分页，每一年有一个目录页，共 1,251 篇，收录至 2026 年 10 月 9 日。2026 年目录页最后更新：2026-10-10；其余年份：2026-08-25。')
table(['年份', '网址', '篇数'], [(y, '/news/%s/' % y, str(n)) for y, n in NEWS] + [('合计', '', '1,251')])

doc.add_heading('四、连到站外的入口', level=2)
table(['名称', '出现在', '连到'], [
    ('线上报名', '首页、招生专区', 'eSchool 新生注册 nanhwa.eschool.edu.my/newreg/'),
    ('WhatsApp 预约参观／咨询', '首页、招生专区', 'wa.me/601164646267'),
    ('校刊书架', '南华数码馆（第四展厅）', 'nanhwappd.github.io/kanwu/'),
    ('Facebook 官方专页', '首页底部', 'facebook.com/sekolahtingginanhwa'),
    ('电话、电邮、地图', '首页底部、招生专区', '05-691 1442／05-691 3457、stnh@nanhwafs.edu.my、Google 地图'),
])
doc.save(out)
print('ok')
